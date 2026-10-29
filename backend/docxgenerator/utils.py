from io import BytesIO
from pathlib import Path
from typing import Dict, Any, Iterable

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_PARAGRAPH_ALIGNMENT

from usuario.models import Usuario


def _set_table_borders(table):
    """Define bordas sólidas na tabela (todas as arestas e internas)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl._element.insert(0, tblPr)

    # remove bordas existentes (se houver)
    for child in list(tblPr):
        if child.tag == qn('w:tblBorders'):
            tblPr.remove(child)

    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        elem = OxmlElement(f'w:{edge}')
        elem.set(qn('w:val'), 'single')   # linha simples
        elem.set(qn('w:sz'), '8')         # espessura
        elem.set(qn('w:color'), '000000') # cor preta
        elem.set(qn('w:space'), '0')
        borders.append(elem)

    tblPr.append(borders)


def _set_table_cell_margins(table, top=80, bottom=80, left=120, right=120):
    """
    Define margens internas (padding) das células em twips (1 pt = 20 twips).
    top/bottom 80 ≈ 4 pt; left/right 120 ≈ 6 pt.
    """
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl._element.insert(0, tblPr)

    # remove configuração anterior (se existir)
    for child in list(tblPr):
        if child.tag == qn('w:tblCellMar'):
            tblPr.remove(child)

    cellMar = OxmlElement('w:tblCellMar')

    top_el = OxmlElement('w:top')
    top_el.set(qn('w:w'), str(top))
    top_el.set(qn('w:type'), 'dxa')

    bottom_el = OxmlElement('w:bottom')
    bottom_el.set(qn('w:w'), str(bottom))
    bottom_el.set(qn('w:type'), 'dxa')

    left_el = OxmlElement('w:left')
    left_el.set(qn('w:w'), str(left))
    left_el.set(qn('w:type'), 'dxa')

    right_el = OxmlElement('w:right')
    right_el.set(qn('w:w'), str(right))
    right_el.set(qn('w:type'), 'dxa')

    cellMar.append(top_el)
    cellMar.append(left_el)
    cellMar.append(bottom_el)
    cellMar.append(right_el)

    tblPr.append(cellMar)


def gerar_docx_plano_estruturado(plano_estruturado: Dict[str, Any], titulo: str = "PLANO DE AULA") -> BytesIO:
    """
    Gera um documento DOCX com layout moderno baseado no plano estruturado.
    """
    doc = Document()
    
    # Configuração da página
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

    # Título principal
    titulo_paragraph = doc.add_paragraph()
    titulo_run = titulo_paragraph.add_run(plano_estruturado.get("titulo", titulo))
    titulo_run.font.name = "Arial"
    titulo_run.font.size = Pt(16)
    titulo_run.bold = True
    titulo_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    titulo_paragraph.space_after = Pt(12)

    # Informações básicas
    info_style = doc.styles['Normal']
    info_style.font.name = "Arial"
    info_style.font.size = Pt(11)
    
    # Duração
    duracao_p = doc.add_paragraph()
    duracao_p.add_run("Duração: ").bold = True
    duracao_p.add_run(plano_estruturado.get("duracao", "50 minutos"))
    duracao_p.space_after = Pt(6)

    # Metodologia
    metodologia_p = doc.add_paragraph()
    metodologia_p.add_run("Metodologia: ").bold = True
    metodologia_p.add_run(plano_estruturado.get("metodologia", "Aprendizagem Ativa"))
    metodologia_p.space_after = Pt(6)

    # Objetivo Geral
    obj_geral_p = doc.add_paragraph()
    obj_geral_p.add_run("Objetivo Geral: ").bold = True
    obj_geral_p.add_run(plano_estruturado.get("objetivo_geral", ""))
    obj_geral_p.space_after = Pt(6)

    # Objetivos Específicos
    obj_esp_p = doc.add_paragraph()
    obj_esp_p.add_run("Objetivos Específicos: ").bold = True
    obj_esp_p.space_after = Pt(3)
    
    objetivos = plano_estruturado.get("objetivos_especificos", [])
    for obj in objetivos:
        obj_item_p = doc.add_paragraph(f"• {obj}")
        obj_item_p.style = "List Bullet"
        obj_item_p.space_after = Pt(3)

    # Recursos
    recursos_p = doc.add_paragraph()
    recursos_p.add_run("Recursos: ").bold = True
    recursos_p.space_after = Pt(3)
    
    recursos = plano_estruturado.get("recursos", [])
    for recurso in recursos:
        recurso_item_p = doc.add_paragraph(f"• {recurso}")
        recurso_item_p.style = "List Bullet"
        recurso_item_p.space_after = Pt(3)

    # Desenvolvimento da Aula
    desenvolvimento = plano_estruturado.get("desenvolvimento", {})
    
    dev_p = doc.add_paragraph()
    dev_p.add_run("Desenvolvimento da Aula: ").bold = True
    dev_p.space_after = Pt(6)

    # Seções do desenvolvimento
    secoes = [
        ("introducao", "Introdução e Revisão"),
        ("desenvolvimento_1", "Desenvolvimento 1"),
        ("desenvolvimento_2", "Desenvolvimento 2"),
        ("aplicacao", "Resolução de Problemas"),
        ("conclusao", "Conclusão e Exercícios")
    ]

    for secao_key, secao_titulo in secoes:
        secao_data = desenvolvimento.get(secao_key, {})
        if not secao_data:
            continue
            
        # Título da seção com tempo
        tempo = secao_data.get("tempo", "")
        titulo_secao = secao_data.get("titulo", secao_titulo)
        
        secao_p = doc.add_paragraph()
        secao_p.add_run(f"({tempo}) {titulo_secao}").bold = True
        secao_p.space_after = Pt(6)

        # Atividades da seção
        atividades = secao_data.get("atividades", [])
        for i, atividade in enumerate(atividades, 1):
            if secao_key == "introducao":
                # Para introdução, usar numeração sequencial
                ativ_p = doc.add_paragraph(f"{i}. {atividade}")
                ativ_p.space_after = Pt(3)
            else:
                # Para outras seções, usar bullet points
                ativ_p = doc.add_paragraph(f"• {atividade}")
                ativ_p.style = "List Bullet"
                ativ_p.space_after = Pt(3)

        # Espaçamento entre seções
        doc.add_paragraph().space_after = Pt(6)

    # Avaliação
    avaliacao_p = doc.add_paragraph()
    avaliacao_p.add_run("Avaliação: ").bold = True
    avaliacao_p.space_after = Pt(6)

    avaliacoes = plano_estruturado.get("avaliacao", [])
    for aval in avaliacoes:
        aval_p = doc.add_paragraph(f"• {aval}")
        aval_p.style = "List Bullet"
        aval_p.space_after = Pt(3)

    # Observações
    obs_p = doc.add_paragraph()
    obs_p.add_run("Observações: ").bold = True
    obs_p.space_after = Pt(6)

    observacoes = plano_estruturado.get("observacoes", [])
    for obs in observacoes:
        obs_item_p = doc.add_paragraph(f"• {obs}")
        obs_item_p.style = "List Bullet"
        obs_item_p.space_after = Pt(3)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def gerar_docx_plano(final_json: Dict[str, Any], titulo: str = "PLANO DE AULA") -> BytesIO:
    doc = Document()

    # ---------------------- Título Centralizado ----------------------
    titulo_paragraph = doc.add_paragraph()
    titulo_paragraph.add_run(titulo.upper())
    titulo_paragraph.alignment = 1  # Centralizado

    # ---------------------- Tabela Principal ----------------------
    col1_width = Inches(2.0)
    col2_width = Inches(4.5)

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.allow_autofit = False
    table.style = "Table Grid"

    # bordas + padding de célula
    _set_table_borders(table)
    _set_table_cell_margins(table, top=80, bottom=80, left=120, right=120)

    def add_row(label: str, value: Any):
        row = table.add_row().cells

        # definir larguras por célula
        row[0].width = col1_width
        row[1].width = col2_width

        # limpar parágrafo vazio padrão
        row[0].paragraphs[0].clear()
        row[1].paragraphs[0].clear()

        row[0].text = label

        if isinstance(value, Iterable) and not isinstance(value, (str, bytes)):
            for idx, item in enumerate(value):
                if idx == 0:
                    p = row[1].paragraphs[0]
                    p.add_run(str(item))
                    p.style = "List Bullet"
                else:
                    row[1].add_paragraph(str(item), style="List Bullet")
        else:
            row[1].text = str(value) if value else ""

    # ---------------------- Seções na ordem do modelo ----------------------
    add_row("Sequência:", final_json.get("Sequência", ""))
    add_row("Ano / Turma", final_json.get("Ano / Turma", ""))
    add_row("Componentes Curriculares", final_json.get("Componentes Curriculares", ""))
    add_row("Competências Gerais", final_json.get("Competências gerais", ""))
    add_row("Competências Específicas", final_json.get("Competências específicas", ""))
    add_row("Objetos de Conhecimento", final_json.get("Objeto de conhecimento", ""))
    add_row("Habilidades Relacionadas", final_json.get("Habilidade BNCC", ""))
    add_row("Metodologia", final_json.get("Tipo da metodologia e descrição da metodologia", ""))
    add_row("Objetivos", final_json.get("Objetivos", ""))
    add_row("Recursos Didáticos", final_json.get("Recursos didáticos", ""))
    add_row("Avaliação", final_json.get("Avaliação", ""))
    add_row("Referências", final_json.get("Referências bibliográficas", ""))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def gerar_docx_resumo(data: Dict[str, Any], usuario: Usuario) -> BytesIO:
    doc = Document()
    sec = doc.sections[0]
    page_width = sec.page_width - sec.left_margin - sec.right_margin

    # ---------- cabeçalho ----------
    hdr = sec.header
    hdr.add_paragraph(f"Prof. {getattr(usuario, 'nome', '')}")
    try:
        hdr.add_table(1, 1, width=page_width)
    except TypeError:  # python-docx < 1.0 fallback
        tbl = hdr.add_table(1, 1)
        tbl.autofit = False
        tbl.columns[0].width = page_width

    # ---------- título ----------
    doc.add_heading(data["title"], level=1)

    # ---------- introdução ----------
    intro = data.get("introduction", "").strip()
    if intro:
        doc.add_paragraph(intro)

    # ---------- seções ----------
    for sec_data in data["sections"]:
        # heading em negrito
        p_head = doc.add_paragraph(sec_data["heading"])
        for run in p_head.runs:
            run.bold = True

        # descrição
        desc = sec_data.get("description", "").strip()
        if desc:
            doc.add_paragraph(desc)

        # bullets
        for bullet in sec_data["bullets"]:
            if isinstance(bullet, str):
                doc.add_paragraph(bullet, style="List Bullet")
            elif isinstance(bullet, dict):
                doc.add_paragraph(bullet.get("question", "Questão"), style="List Bullet")
                for step in bullet.get("solution", []):
                    doc.add_paragraph(step, style="List Bullet 2")

    # ---------- conclusão ----------
    concl = doc.add_paragraph(data["conclusion"])
    for run in concl.runs:
        run.bold = True

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
