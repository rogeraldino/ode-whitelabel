import io
import json
import logging
from typing import List, Tuple

from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated

from docx import Document
from docx.shared import Inches, Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import qrcode

from atividades.models import ListaAtividade, Questao, Alternativa
from usuario.permissions import perfil_perm

logger = logging.getLogger(__name__)


# ========= Utilitários XML / Estilo =========

def _twips_from_mm(mm_val: float) -> int:
    return int((mm_val / 25.4) * 1440)

def _set_cell_bottom_border(cell, sz=8, color='000000', val='single'):
    """Borda inferior (linha de preenchimento)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn('w:tcBorders'):
            tcPr.remove(child)
    borders = OxmlElement('w:tcBorders')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), val)
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), color)
    borders.append(bottom)
    tcPr.append(borders)

def _clear_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn('w:tcBorders'):
            tcPr.remove(child)

def _cell_clear_paragraphs(cell):
    for p in cell.paragraphs:
        for r in p.runs:
            r.clear()
        p.clear()

def _add_text(cell, text, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, underline=False):
    _cell_clear_paragraphs(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.underline = underline
    p.alignment = align
    return p, run

def _add_center(cell, text, **kw):
    return _add_text(cell, text, align=WD_ALIGN_PARAGRAPH.CENTER, **kw)

def _qr_bytes_from_text(text: str) -> io.BytesIO:
    img = qrcode.make(text)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf

def _set_table_width_exact(table, width_mm: float):
    """Força a largura total da tabela (não só da célula)."""
    tbl = table._tbl
    tblPr = tbl.tblPr or OxmlElement('w:tblPr')
    if tbl.tblPr is None:
        tbl._element.insert(0, tblPr)
    # remove w:tblW prévio
    for child in list(tblPr):
        if child.tag == qn('w:tblW'):
            tblPr.remove(child)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(_twips_from_mm(width_mm)))  # em twips
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)

def _table_outer_border(table, sz=8, color='000000'):
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        table._tbl._element.insert(0, tblPr)
    # remove bordas existentes
    for child in list(tblPr):
        if child.tag == qn('w:tblBorders'):
            tblPr.remove(child)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        elem = OxmlElement(f'w:{edge}')
        elem.set(qn('w:val'), 'single')
        elem.set(qn('w:sz'), str(sz))
        elem.set(qn('w:space'), '0')
        elem.set(qn('w:color'), color)
        borders.append(elem)
    # sem insideH/insideV → só a caixa externa
    tblPr.append(borders)


# ========= Cabeçalho + BNCC + Nota (separada) =========

def _adicionar_cabecalho(doc: Document, lista: ListaAtividade):
    """
    1) BNCC no topo direito (sublinhado).
    2) Tabela esquerda com labels (30%) e linhas (70%),
       com 'Data' em linha própria SEM linha completa (apenas '__/____/____').
    3) Caixa 'Nota' separada, CENTRALIZADA, com largura EXATA de 50 mm (igual ao QR).
    """
    # 1) BNCC topo direito
    codigos = sorted({h.codigo_bncc for h in lista.habilidades.all() if getattr(h, 'codigo_bncc', None)})
    if codigos:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("Códigos BNCC: " + ", ".join(codigos))
        r.font.size = Pt(9)
        r.underline = True

    # 2) Cabeçalho (labels + linhas) — 5 linhas (inclui Data)
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    # Proporção 30% / 70% (considerando ~5.7" de largura útil)
    table.columns[0].width = Inches(1.7)   # ~30%
    table.columns[1].width = Inches(4.0)   # ~70%

    labels = ["Professor:", "Componente:", "Estudante:", "Turma:", "Data:"]
    for i, label in enumerate(labels):
        table.rows[i].cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        table.rows[i].cells[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        _add_text(table.rows[i].cells[0], label, size=10)
        _add_text(table.rows[i].cells[1], " ", size=10)
        _clear_cell_borders(table.rows[i].cells[0])

        # Borda inferior apenas nas 4 primeiras (Professor/Componente/Estudante/Turma).
        if i < 4:
            _set_cell_bottom_border(table.rows[i].cells[1], sz=8)
        else:
            # Linha 'Data' sem underline completo — apenas o placeholder:
            _clear_cell_borders(table.rows[i].cells[1])
            data_para = table.rows[i].cells[1].paragraphs[0]
            data_para.add_run("__/____/____").font.size = Pt(10)

    # respiro
    doc.add_paragraph("")

    # 3) Caixa Nota — largura exata 50mm, centralizada
    nota_tbl = doc.add_table(rows=2, cols=1)
    nota_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    nota_tbl.autofit = False
    nota_tbl.columns[0].width = Mm(50)
    _set_table_width_exact(nota_tbl, 50)   # força a tabela a ter 50mm de largura total
    _table_outer_border(nota_tbl, sz=8)
    nota_tbl.cell(0, 0).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _add_center(nota_tbl.rows[0].cells[0], "Nota", bold=True, size=12)
    _add_center(nota_tbl.rows[1].cells[0], " ", size=12)

    # espaçamento depois da nota
    doc.add_paragraph("")


# ========= QRCode central =========

def _adicionar_qrcode_central(doc: Document, lista: ListaAtividade):
    """Apenas o QRCode centralizado (50 mm)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr_buf = _qr_bytes_from_text(str(lista.uuid))
    p.add_run().add_picture(qr_buf, width=Mm(50))
    doc.add_paragraph("")


# ========= Questões (coluna única) =========

def _add_discursive_lines_table(doc: Document, count: int = 4):
    """
    Gera uma tabela count×1 com borda inferior em cada linha.
    """
    t = doc.add_table(rows=count, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = True
    for i in range(count):
        cell = t.rows[i].cells[0]
        _clear_cell_borders(cell)
        _set_cell_bottom_border(cell, sz=8)
        _add_text(cell, " ", size=9)
    doc.add_paragraph("")

def _adicionar_questoes(doc: Document, questoes: List[Tuple[Questao, float]]):
    for idx, (q, peso) in enumerate(questoes, start=1):
        # Título da questão
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(f"Q.{idx} ({peso:.1f}) ")
        r1.bold = True
        r1.font.size = Pt(11)
        r2 = p.add_run("- " + (q.enunciado or ""))
        r2.font.size = Pt(11)

        if q.tipo == "objetiva":
            for alt in Alternativa.objects.filter(questao=q).order_by("letra"):
                pa = doc.add_paragraph()
                pa.paragraph_format.space_before = Pt(0)
                pa.paragraph_format.space_after = Pt(2)
                ra = pa.add_run(f"{alt.letra}) (   ) ")
                ra.bold = True
                ra.font.size = Pt(10)
                rb = pa.add_run(alt.texto or "")
                rb.font.size = Pt(10)
        else:
            _add_discursive_lines_table(doc, count=4)

        doc.add_paragraph("")


# ========= View principal =========

@api_view(["POST"])
@permission_classes([IsAuthenticated, perfil_perm("corrigir-atividade")])
def gerar_correcao_atividade_docx(request):
    """
    Payload:
    {
        "lista_atividade_id": "<uuid da ListaAtividade>",
        "questoes": {
            "<questao_uuid_1>": 1,
            "<questao_uuid_2>": 0.5,
            ...
        }
    }
    """
    try:
        payload = request.data if isinstance(request.data, dict) else json.loads(request.body.decode("utf-8"))

        lista_uuid = payload.get("lista_atividade_id")
        qmap = payload.get("questoes", {})

        if not lista_uuid or not isinstance(qmap, dict) or not qmap:
            return JsonResponse({"error": "Payload inválido. Envie lista_atividade_id e questoes."}, status=400)

        lista = get_object_or_404(ListaAtividade, uuid=lista_uuid)

        # manter ordem recebida
        questoes: List[Tuple[Questao, float]] = []
        for q_uuid, peso_raw in qmap.items():
            q = get_object_or_404(Questao, uuid=q_uuid)
            try:
                peso = float(peso_raw)
            except Exception:
                peso = float((peso_raw or {}).get("peso", 1.0)) if isinstance(peso_raw, dict) else 1.0
            questoes.append((q, peso))

        doc = Document()

        _adicionar_cabecalho(doc, lista)
        # _adicionar_qrcode_central(doc, lista)
        _adicionar_questoes(doc, questoes)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        resp = HttpResponse(
            buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        resp["Content-Disposition"] = f'attachment; filename="{(lista.nome or "lista")}-correcao.docx"'
        logger.info("DOCX gerado (lista=%s, user=%s)", lista.uuid, getattr(request.user, 'id', None))
        return resp

    except Exception:
        logger.exception("Erro ao gerar DOCX da lista")
        return HttpResponse("Erro ao gerar documento.", status=500)
